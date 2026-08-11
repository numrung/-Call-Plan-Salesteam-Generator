import io
import gc
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

def process_call_plan_logic(df, job_number):
    """
    Stateless Function: ไม่เก็บ State/Cache ป้องกันข้อมูลปนกันระหว่างผู้ใช้
    """
    # 1. คัดลอก DataFrame เพื่อป้องกันการแก้ไข DataFrame ต้นฉบับที่อาจแชร์ร่วมกัน
    local_df = df.copy()
    output_bundles = {}

    col_filter_idx = 9
    col_filter = local_df.columns[col_filter_idx]
    
    # ลบแถว Header ซ้ำ
    local_df = local_df[~local_df[col_filter].astype(str).str.contains("รหัสพนักงาน|^\s*$", na=False)]

    columns = local_df.columns.tolist()
    col_box, col_store, col_emp_id = columns[4], columns[5], columns[6]
    date_columns = columns[12:43]

    emp_ids = local_df[col_filter].dropna().astype(str).str.strip()
    emp_ids = emp_ids[emp_ids != ""].unique()

    for emp_id in emp_ids:
        result_rows = []
        emp_data = local_df[local_df[col_filter].astype(str).str.strip() == emp_id]

        for i, col_date in enumerate(date_columns):
            matched_rows = emp_data[emp_data[col_date] == 1]
            for _, row in matched_rows.iterrows():
                result_rows.append([
                    row[col_box],
                    row[col_store],
                    row[col_emp_id],
                    i + 1
                ])

        if result_rows:
            temp_df = pd.DataFrame(result_rows, columns=["กล่อง", "รหัสร้านค้า", "รหัสพนักงาน", "วันที่"])
            temp_df.sort_values(by="วันที่", inplace=True)
            temp_df["ลำดับ"] = temp_df.groupby("วันที่").cumcount() + 1
            output_df = temp_df[["ลำดับ", "กล่อง", "รหัสร้านค้า", "รหัสพนักงาน", "วันที่"]]

            # ใช้ context manager สร้าง BytesIO ชั่วคราวเฉพาะรอบลูป
            with io.BytesIO() as final_out:
                with pd.ExcelWriter(final_out, engine='openpyxl') as writer:
                    output_df.to_excel(writer, index=False, sheet_name='Data')
                    ws = writer.sheets['Data']
                    
                    max_row, max_col = ws.max_row, ws.max_column
                    col_letter = get_column_letter(max_col)
                    
                    clean_emp_id = "".join(e for e in emp_id if e.isalnum() or e == "_")
                    tab = Table(displayName=f"Table_{clean_emp_id}", ref=f"A1:{col_letter}{max_row}")
                    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, 
                                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
                    tab.tableStyleInfo = style
                    ws.add_table(tab)
                
                filename = f"Call Plan_{job_number}_{emp_id}.xlsx"
                # ดึงเฉพาะค่า raw bytes ไปเก็บไว้
                output_bundles[filename] = final_out.getvalue()

    # 2. บังคับคืนหน่วยความจำและลบตัวแปรชั่วคราว
    del local_df
    gc.collect()

    return True, output_bundles
